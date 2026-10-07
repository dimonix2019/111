#!/usr/bin/env python3
"""Сборка резюме: ведущий аналитик процессов и BI."""

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips

OUT = "/workspace/Abramov_DN_resume_ANALYST_process_BI.docx"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
DARK = RGBColor(0x22, 0x22, 0x22)
GRAY = RGBColor(0x55, 0x55, 0x55)


def set_run(run, size, bold=False, color=DARK, name="Arial"):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), name)


def add_p(doc, text, size=10, bold=False, color=DARK, space_before=0, space_after=40, bullet=False):
    style = "List Bullet" if bullet else "Normal"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if bullet:
        p.paragraph_format.left_indent = Twips(360)
    run = p.add_run(text)
    set_run(run, size, bold=bold, color=color)
    return p


def section(doc, title):
    add_p(doc, title, size=13, bold=True, color=NAVY, space_before=12, space_after=6)


def job_dates(doc, text):
    add_p(doc, text, size=10, color=GRAY, space_before=10, space_after=0)


def company(doc, text):
    add_p(doc, text, size=11, bold=True, color=NAVY, space_before=2, space_after=0)


def meta(doc, text):
    add_p(doc, text, size=9.5, color=GRAY, space_before=0, space_after=0)


def role(doc, text):
    add_p(doc, text, size=10.5, bold=True, color=DARK, space_before=4, space_after=4)


def label(doc, text):
    add_p(doc, text, size=10, bold=True, color=DARK, space_before=4, space_after=2)


def bullet(doc, text):
    add_p(doc, text, size=10, bullet=True, space_before=0, space_after=2)


def course(doc, title, place):
    add_p(doc, title, size=10, space_before=4, space_after=0)
    add_p(doc, place, size=9.5, color=GRAY, space_before=0, space_after=2)


def main():
    doc = Document()
    section_page = doc.sections[0]
    section_page.page_width = Twips(11906)
    section_page.page_height = Twips(16838)
    section_page.left_margin = Twips(907)
    section_page.right_margin = Twips(907)
    section_page.top_margin = Twips(680)
    section_page.bottom_margin = Twips(680)

    add_p(doc, "Абрамов Дмитрий Николаевич", size=16, bold=True, color=NAVY, space_after=2)
    add_p(doc, "Мужчина", size=10, color=GRAY, space_after=2)
    add_p(doc, "+7 (903) 155-89-49", size=10.5, space_after=0)
    add_p(doc, "abramov-dn@yandex.ru - предпочитаемый способ связи", size=10.5, space_after=2)
    add_p(doc, "Проживает: Москва", size=10, color=GRAY, space_after=0)
    add_p(doc, "Гражданство: Россия, есть разрешение на работу: Россия", size=10, color=GRAY, space_after=0)
    add_p(doc, "Не готов к переезду, готов к редким командировкам", size=10, color=GRAY, space_after=6)

    section(doc, "Желаемая должность и зарплата")
    add_p(doc, "Ведущий аналитик процессов и BI (мониторинг, дашборды)", size=12, bold=True, space_after=4)
    add_p(doc, "Специализации:", size=10, space_after=0)
    add_p(doc, "- Ведущий аналитик", size=10, space_after=0)
    add_p(doc, "- Аналитик", size=10, space_after=4)
    add_p(doc, "Тип занятости: полная занятость", size=10, space_after=0)
    add_p(doc, "Формат работы: на месте работодателя, удалённо, гибрид", size=10, space_after=0)
    add_p(doc, "Желательное время в пути до работы: не имеет значения", size=10, space_after=2)
    add_p(doc, "400 000 руб. на руки", size=11, bold=True, space_after=8)

    label(doc, "Обо мне")
    add_p(
        doc,
        "Строю процесс, ставлю его на мониторинг и подключаю BI. В Сбере с нуля запустил CSI по департаменту и сквозной контроль работы с инициативами: правила, статусы, регулярная отчётность, дашборд.",
        size=10.5,
        space_after=4,
    )
    add_p(
        doc,
        "В Axenix собирал управленческую отчётность для руководства. В ПСБ веду хранилище, чтобы такие контуры стояли на данных. Ищу роль ведущего аналитика по процессам, мониторингу и дашбордам.",
        size=10.5,
        space_after=8,
    )

    section(doc, "Опыт работы - 24 года 9 месяцев")

    job_dates(doc, "Ноябрь 2024 - настоящее время    1 год 11 месяцев")
    company(doc, "Банк ПСБ")
    meta(doc, "Москва, job.psbank.ru")
    meta(doc, "Финансовый сектор")
    meta(doc, "• Банк")
    role(doc, "Технический руководитель проекта хранилища данных")
    label(doc, "Достижения и результаты:")
    bullet(
        doc,
        "Собрал хранилище на открытом стеке: контуры загрузки, витрины, согласование с архитектурой и информационной безопасностью. Дальше - подключение BI.",
    )
    bullet(
        doc,
        "Выстроил поставку изменений по контурам хранилища, довёл решение до промышленного контура.",
    )
    bullet(
        doc,
        "Согласовал цели, дорожную карту, бизнес-требования и ТЗ с ИТ, ИБ и бизнесом. Держу понятный статус по работам.",
    )
    label(doc, "Обязанности:")
    bullet(doc, "Контуры данных, витрины, подготовка к BI.")
    bullet(doc, "Согласование решений с архитектурой и информационной безопасностью.")
    bullet(doc, "Работа с бизнесом, ИТ и подрядчиками по составу работ и срокам.")

    job_dates(doc, "Сентябрь 2022 - Ноябрь 2024    2 года 3 месяца")
    company(doc, "Axenix (ex Accenture)")
    meta(doc, "Москва, axenix.pro")
    role(doc, "Консультант по управленческой отчётности / поставка BI")
    label(doc, "Достижения и результаты:")
    bullet(
        doc,
        "Собрал систему управленческой отчётности и BI для топ-менеджмента НЛМК на платформе Форсайт: дашборды, KPI, регулярная аналитика. В части экранов использовали модели для оценки и прогноза.",
    )
    bullet(
        doc,
        "Доводил работу от задачи заказчика до метрик, экрана, приёмки и правок по обратной связи.",
    )
    label(doc, "Обязанности:")
    bullet(doc, "Разбор задачи, состав метрик и экранов, постановка, приёмка.")
    bullet(doc, "Работа с CIO и бизнес-заказчиками по управленческой отчётности.")

    job_dates(doc, "Август 2016 - Сентябрь 2022    6 лет 2 месяца")
    company(doc, "Сбер")
    meta(doc, "Москва, rabota.sber.ru")
    meta(doc, "Финансовый сектор")
    meta(doc, "• Банк")
    role(doc, "Эксперт по отчётности и BI, вёл контуры аналитики (SberData)")
    label(doc, "Достижения:")
    bullet(
        doc,
        "С нуля собрал процесс измерения CSI по департаменту: методика, сбор данных, дашборды на Qlik Sense и QlikView в разрезе продуктов, клиентов (блоки, ДИТ) и динамики.",
    )
    bullet(
        doc,
        "Выстроил процесс работы с бизнес- и ИТ-инициативами и поставил его на мониторинг. Сделал на QlikView дашборд сквозного контроля, в том числе по инициативам с ИИ.",
    )
    bullet(
        doc,
        "Вместе с департаментом корпоративной архитектуры разобрал модули инициатив, согласование архитектуры и раздачу задач продуктовым командам.",
    )
    bullet(
        doc,
        "Связал процесс инициатив с управлением ресурсами. Собрал Portfolio Market Place: планирование задач команд с учётом зависимостей.",
    )
    bullet(
        doc,
        "Собрал базу знаний программы «Создание Фабрики данных» и дашборд по ней: ресурсы, инициативы, коллегиальные органы. Настроил сквозной поиск и аналитику.",
    )
    bullet(
        doc,
        "Поставил регулярную отчётность по работе с инициативами на QlikView.",
    )
    label(doc, "Обязанности:")
    bullet(doc, "Сбор потребности, согласование метрик и варианта реализации.")
    bullet(doc, "Ведение контуров отчётности и мониторинга: бэклог, дашборды, приёмка, обратная связь.")
    bullet(doc, "Связка процесса инициатив со смежными контурами и статусы для руководства.")

    job_dates(doc, "Декабрь 2011 - Август 2016    4 года 9 месяцев")
    company(doc, "Сбер")
    meta(doc, "Москва, rabota.sber.ru")
    meta(doc, "Финансовый сектор")
    meta(doc, "• Банк")
    role(doc, "Руководитель направления в дирекции программы «Централизация 2.0» (единая ИТ-платформа банка)")
    label(doc, "Достижения:")
    bullet(doc, "Сделал дашборд внедрений программы на QlikView и дашборд по коммуникациям.")
    bullet(doc, "Собрал процесс коммуникаций программы: стейкхолдеры, внутреннее и внешнее окружение, материалы для УК и Правления.")
    bullet(
        doc,
        "Программа заняла 1 место IPMA International Project Excellence Award 2015 (крупные проекты) и победила в Celent Model Bank 2016. Я вёл отчётность, коммуникации и материалы для коллегиальных органов.",
    )
    label(doc, "Обязанности:")
    bullet(doc, "Отчёты для руководства банка и контролирующих органов.")
    bullet(doc, "Сбор данных от команд, процесс коммуникаций, материалы к комитетам.")

    job_dates(doc, "Февраль 2010 - Декабрь 2011    1 год 11 месяцев")
    role(doc, "ЗАО «Беллинтегратор». Менеджер проектов BI и мониторинга")
    add_p(doc, "Больше 30 проектов BI и мониторинга: МТС, ВымпелКом, МегаФон, банки, ЦППК.", size=10, space_after=2)
    add_p(doc, "Требования, БТ и ТЗ, источники данных, постановка, приёмка, обучение пользователей.", size=10, space_after=2)
    add_p(doc, "Подбор решения BI или мониторинга под задачу заказчика.", size=10, space_after=2)

    job_dates(doc, "Декабрь 2007 - Декабрь 2009    2 года 1 месяц")
    role(doc, "ЗАО «Верисел Проекты». Менеджер проектов по системам мониторинга")
    add_p(doc, "Пилот мониторинга в ЦентрТелекоме, проекты МТС и ЦентрТелекома.", size=10, space_after=2)
    add_p(doc, "Мониторинг бизнес-приложений в ИНГ Банк (Евразия).", size=10, space_after=2)
    add_p(doc, "Проектная документация и координация субподрядчиков.", size=10, space_after=2)

    job_dates(doc, "Январь 2007 - Декабрь 2007    1 год")
    role(doc, "МТС. Руководитель группы развития системы мониторинга")
    add_p(doc, "Архитектура и прототип единой системы мониторинга: данные и экраны бизнес-KPI.", size=10, space_after=2)
    add_p(doc, "Регламенты и инструкции. Защита бюджета на инвесткомитете.", size=10, space_after=2)
    add_p(doc, "Постановка задач интеграторам, поддержка пользователей.", size=10, space_after=2)

    job_dates(doc, "Январь 2002 - Январь 2007    5 лет 1 месяц")
    company(doc, "ОАО «Транстелеком»")
    meta(doc, "www.transtelecom.ru")
    meta(doc, "Провайдер услуг интернета и телефонии")
    role(doc, "Программист, инженер биллинга")
    bullet(doc, "Собрал систему биллинга (C#, SQL Server).")
    bullet(doc, "Собрал CRM для отдела продаж (C#, SQL Server, интерфейс из метаданных).")
    bullet(doc, "Тарифы, счета, поддержка биллинга, CRM, домена и почты.")

    section(doc, "Образование")
    add_p(doc, "Высшее", size=10, space_after=0)
    add_p(doc, "2001", size=10, color=GRAY, space_after=0)
    add_p(
        doc,
        "Московский государственный технический университет радиотехники, электроники и автоматики, Москва",
        size=10.5,
        bold=True,
        space_after=0,
    )
    add_p(doc, "Кибернетики, Искусственный интеллект", size=10, space_after=6)

    section(doc, "Повышение квалификации, курсы")
    course(doc, "2025  Аренадата. Архитектура DWH", "Аренадата, Архитектура DWH")
    course(doc, "2023  Профессия «Аналитик»", "Институт бизнес-аналитики - подготовка данных, построение дашбордов")
    course(doc, "2021  Qlik Sense. Внутренние курсы", "Sber")
    course(doc, "2020  Управление продуктами по Agile", "Сбер")
    course(doc, "2016  «Зелёный пояс» - оптимизация бизнес-процессов", "ПАО «Сбербанк»")
    course(doc, "2015  Управление проектами - IPMA", "СОВНЕТ")
    course(doc, "2015  «Подготовка презентаций»", "Авторский курс Алексея Каптерева")
    course(doc, "2014  «Мастерство переговоров»", "Университет риторики и ораторского мастерства")
    course(doc, "2011  QlikTech", "Вендорские курсы по QlikView")

    section(doc, "Тесты, экзамены")
    course(doc, "2019  Управленческий анализ и Большие данные (Big Data II)", "ПАО «Сбербанк»")
    course(doc, "2014  Microsoft Certified Professional (MCP)", "Разработка приложений для Windows на C#")
    course(doc, "2011  Разработка приложений на QlikView", "Курсы у вендора QlikTech")
    course(doc, "2011  Quest Software", "Сертификат по решениям для управления и мониторинга баз данных")

    section(doc, "Навыки")
    add_p(doc, "Знание языков", size=10, bold=True, space_after=2)
    add_p(doc, "Русский - Родной", size=10, space_after=0)
    add_p(doc, "Английский - B2 - Средне-продвинутый", size=10, space_after=6)
    add_p(doc, "Ключевые навыки", size=10, bold=True, space_after=2)
    add_p(
        doc,
        "Описание процессов  ·  Мониторинг процессов  ·  CSI  ·  Управление инициативами  ·  BI  ·  Дашборды  ·  QlikView  ·  Qlik Sense  ·  Управленческая отчётность  ·  KPI / SLA  ·  Постановка требований  ·  Приёмка  ·  BPMN  ·  Хранилище данных  ·  Витрины  ·  SQL  ·  MS Excel  ·  PowerPoint  ·  Jira  ·  Confluence  ·  Agile",
        size=10,
        space_after=8,
    )

    section(doc, "Опыт вождения")
    add_p(doc, "Имеется собственный автомобиль", size=10, space_after=0)
    add_p(doc, "Права категории C", size=10, space_after=8)

    section(doc, "Дополнительная информация")
    label(doc, "Обо мне")
    add_p(
        doc,
        "В Сбере сам предлагал и собирал дашборды: CSI и мониторинг инициатив. В Axenix - управленческая отчётность. В ПСБ - хранилище и витрины под такие контуры.",
        size=10.5,
        space_after=4,
    )
    add_p(
        doc,
        "Привык доводить работу до экрана, которым пользуются, а не до описания процесса в папке.",
        size=10.5,
        space_after=2,
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
