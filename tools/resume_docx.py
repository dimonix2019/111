"""Общая вёрстка резюме под формат hh."""

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
DARK = RGBColor(0x22, 0x22, 0x22)
GRAY = RGBColor(0x55, 0x55, 0x55)


def set_run(run, size, bold=False, color=DARK, name="Arial"):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), name)


def add_p(doc, text, size=10, bold=False, color=DARK, space_before=0, space_after=40, bullet=False):
    style = "List Bullet" if bullet else "Normal"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if bullet:
        p.paragraph_format.left_indent = Twips(360)
    run = p.add_run(text)
    set_run(run, size, bold=bold, color=color)
    return p


def section(doc, title):
    add_p(doc, title, size=13, bold=True, color=NAVY, space_before=12, space_after=6)


def job_dates(doc, text):
    add_p(doc, text, size=10, color=GRAY, space_before=10, space_after=0)


def company(doc, text):
    add_p(doc, text, size=11, bold=True, color=NAVY, space_before=2, space_after=0)


def meta(doc, text):
    add_p(doc, text, size=9.5, color=GRAY, space_before=0, space_after=0)


def role(doc, text):
    add_p(doc, text, size=10.5, bold=True, color=DARK, space_before=4, space_after=4)


def label(doc, text):
    add_p(doc, text, size=10, bold=True, color=DARK, space_before=4, space_after=2)


def bullet(doc, text):
    add_p(doc, text, size=10, bullet=True, space_before=0, space_after=2)


def course(doc, title, place):
    add_p(doc, title, size=10, space_before=4, space_after=0)
    add_p(doc, place, size=9.5, color=GRAY, space_before=0, space_after=2)


def new_doc():
    doc = Document()
    page = doc.sections[0]
    page.page_width = Twips(11906)
    page.page_height = Twips(16838)
    page.left_margin = Twips(907)
    page.right_margin = Twips(907)
    page.top_margin = Twips(680)
    page.bottom_margin = Twips(680)
    return doc


def header(doc):
    add_p(doc, "Абрамов Дмитрий Николаевич", size=16, bold=True, color=NAVY, space_after=2)
    add_p(doc, "Мужчина", size=10, color=GRAY, space_after=2)
    add_p(doc, "+7 (903) 155-89-49", size=10.5, space_after=0)
    add_p(doc, "abramov-dn@yandex.ru - предпочитаемый способ связи", size=10.5, space_after=2)
    add_p(doc, "Проживает: Москва", size=10, color=GRAY, space_after=0)
    add_p(doc, "Гражданство: Россия, есть разрешение на работу: Россия", size=10, color=GRAY, space_after=0)
    add_p(doc, "Не готов к переезду, готов к редким командировкам", size=10, color=GRAY, space_after=6)


def job_block(doc, dates, company_name, meta_lines, role_title):
    job_dates(doc, dates)
    if company_name:
        company(doc, company_name)
    for line in meta_lines:
        meta(doc, line)
    role(doc, role_title)


def education_block(doc):
    section(doc, "Образование")
    add_p(doc, "Высшее", size=10, space_after=0)
    add_p(doc, "2001", size=10, color=GRAY, space_after=0)
    add_p(
        doc,
        "Московский государственный технический университет радиотехники, электроники и автоматики, Москва",
        size=10.5,
        bold=True,
        space_after=0,
    )
    add_p(doc, "Кибернетики, Искусственный интеллект", size=10, space_after=6)


def courses_block(doc, extra_first=None):
    section(doc, "Повышение квалификации, курсы")
    items = extra_first or []
    items += [
        ("2025  Аренадата. Архитектура DWH", "Аренадата"),
        ("2023  Профессия «Аналитик»", "Институт бизнес-аналитики - подготовка данных, дашборды"),
        ("2021  Qlik Sense. Внутренние курсы", "Сбер"),
        ("2020  Управление продуктами по Agile", "Сбер"),
        ("2016  «Зелёный пояс» - оптимизация бизнес-процессов", "ПАО «Сбербанк»"),
        ("2015  Управление проектами - IPMA", "СОВНЕТ"),
        ("2011  QlikView", "Вендорские курсы QlikTech"),
    ]
    seen = set()
    for title, place in items:
        if title in seen:
            continue
        seen.add(title)
        course(doc, title, place)


def exams_block(doc):
    section(doc, "Тесты, экзамены")
    course(doc, "2019  Управленческий анализ и Большие данные (Big Data II)", "ПАО «Сбербанк»")
    course(doc, "2014  Microsoft Certified Professional (MCP)", "Разработка приложений для Windows на C#")
    course(doc, "2011  Разработка приложений на QlikView", "Курсы у вендора QlikTech")


def languages_and_skills(doc, skills):
    section(doc, "Навыки")
    add_p(doc, "Знание языков", size=10, bold=True, space_after=2)
    add_p(doc, "Русский - Родной", size=10, space_after=0)
    add_p(doc, "Английский - B2 - Средне-продвинутый", size=10, space_after=6)
    add_p(doc, "Ключевые навыки", size=10, bold=True, space_after=2)
    add_p(doc, skills, size=10, space_after=8)


def driving(doc):
    section(doc, "Опыт вождения")
    add_p(doc, "Имеется собственный автомобиль", size=10, space_after=0)
    add_p(doc, "Права категории C", size=10, space_after=8)
